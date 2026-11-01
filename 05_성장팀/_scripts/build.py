# -*- coding: utf-8 -*-
"""팀챌린지_향상점수계산.xlsx 생성 (2026-11-02, 결재2 GPT 교차검토 반영판)
사용: python3 build.py <출력.xlsx> [시나리오이름]   (시나리오는 data.SCENARIOS, 기본 1_확정값)"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.workbook.defined_name import DefinedName
import data

YEL, GRN, GRY, BLU, PNK = "FFF2CC", "E2EFDA", "EDEDED", "DDEBF7", "FCE4D6"
P_FIRST, P_LAST = 3, 52      # 참가자 행
R_LAST = 102                 # 붙여넣기 행 끝
WRAP = Alignment(wrap_text=True, vertical="center")

def fill(c): return PatternFill("solid", fgColor=c)

def header(ws, row, col, text, color):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = fill(color); c.font = Font(bold=True); c.alignment = WRAP

def title(ws, cell, text):
    ws[cell] = text; ws[cell].font = Font(bold=True, size=11)

# ---- 붙여넣기 영역(R~U)과 자동 정리 칸(W~AA): 진단·평가A·B·C 공통 ----
PASTE_COLS = ["R", "S", "T", "U"]
HELP = {"code": "W", "pct": "X", "nc": "Y", "roster": "Z", "cnt": "AA"}

def paste_area(ws, rows, items_name):
    ws["R1"] = "붙여넣기 영역 (노란 칸, 3행부터)"; ws["R1"].font = Font(bold=True)
    for col, t in zip(PASTE_COLS, ["참가자 코드", "맞힌 수", "문항 수", "정답률(%)"]):
        header(ws, 2, ws[col + "1"].column, t, YEL)
    for col, t in zip(["W", "X", "Y", "Z", "AA"],
                      ["정리된 코드(자동)", "계산 정답률(자동)", "비경쟁(자동)", "명단 점검(자동)", "맞힌 수(자동)"]):
        header(ws, 2, ws[col + "1"].column, t, GRY)
    for r in range(3, R_LAST + 1):
        for col in PASTE_COLS:
            ws[f"{col}{r}"].fill = fill(YEL)
        if r - 3 < len(rows):
            for col, v in zip(PASTE_COLS, rows[r - 3]):
                if v is not None:
                    ws[f"{col}{r}"] = v
        u = f'UPPER(TRIM(R{r}&""))'
        ws[f"W{r}"] = f'=IF({u}="","",IF(RIGHT({u},2)="-N",TRIM(LEFT({u},LEN({u})-2)),{u}))'
        ws[f"X{r}"] = (f'=IF(W{r}="","",IF(ISNUMBER(U{r}),ROUND(U{r},1),IF(AND(ISNUMBER(S{r}),ISNUMBER(T{r})),'
                       f'IF(T{r}>0,ROUND(S{r}/T{r}*100,1),""),"")))')
        ws[f"Y{r}"] = f'=IF(W{r}="","",IF(RIGHT({u},2)="-N","비경쟁",""))'
        ws[f"Z{r}"] = f'=IF(W{r}="","",IF(COUNTIF(참가자!$A$3:$A$52,W{r})=0,"명단에 없는 코드",""))'
        ws[f"AA{r}"] = (f'=IF(W{r}="","",IF(ISNUMBER(U{r}),ROUND(U{r}*{items_name}/100,1),'
                        f'IF(AND(ISNUMBER(S{r}),ISNUMBER(T{r})),IF(T{r}>0,S{r},""),"")))')
    for col, w in {"Q": 3, "R": 12, "S": 8, "T": 8, "U": 10, "V": 3, "W": 12, "X": 10, "Y": 10, "Z": 16, "AA": 10}.items():
        ws.column_dimensions[col].width = w

def lookup(col, r):
    return f'IFERROR(INDEX(${col}$3:${col}$102,MATCH(A{r},$W$3:$W$102,0)),"")'

def common_left(ws, r, cnt_col, pct_col, st_col, dup_col):
    ws[f"A{r}"] = f'=IF(참가자!A{r}="","",UPPER(TRIM(참가자!A{r})))'
    ws[f"B{r}"] = f'=IF(A{r}="","",IF(참가자!B{r}="","(팀 없음)",TRIM(참가자!B{r})))'
    ws[f"{cnt_col}{r}"] = f'=IF(A{r}="","",{lookup("AA", r)})'
    ws[f"{pct_col}{r}"] = f'=IF(A{r}="","",{lookup("X", r)})'
    ws[f"{st_col}{r}"] = (f'=IF(A{r}="","",IF({pct_col}{r}="","결석",IF({lookup("Y", r)}="비경쟁","비경쟁 응시","응시")))')
    ws[f"{dup_col}{r}"] = f'=IF(A{r}="","",IF(COUNTIF($W$3:$W$102,A{r})>1,"중복! 확인",""))'

def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

def build(path, overrides=None):
    vals = dict(data.DEFAULTS); vals.update(overrides or {})
    wb = Workbook()
    # ---------------- 참가자 ----------------
    ws = wb.active; ws.title = "참가자"
    title(ws, "A1", "참가자 명단 — 실명·학번·연락처 금지. 무작위 6자리 참가자 코드만(예: K7M2QX). (예시 12명은 가짜 데이터)")
    for i, t in enumerate(["참가자 코드", "팀", "메모(선택, 개인정보 금지)"], 1): header(ws, 2, i, t, GRY)
    for r in range(P_FIRST, P_LAST + 1):
        ws[f"A{r}"].fill = fill(YEL); ws[f"B{r}"].fill = fill(YEL)
        if r - 3 < len(data.PARTICIPANTS):
            ws[f"A{r}"], ws[f"B{r}"] = data.PARTICIPANTS[r - 3]
    widths(ws, {"A": 14, "B": 10, "C": 30}); ws.freeze_panes = "A3"

    # ---------------- 진단 ----------------
    ws = wb.create_sheet("진단")
    title(ws, "A1", "진단 (1회차, 12문항) — 팀 편성 참고 + 수준 파악에만 씀. 점수(기준점수)에는 쓰지 않음")
    for i, t in enumerate(["참가자 코드", "팀", "맞힌 수", "정답률(%)", "상태", "중복 점검"], 1): header(ws, 2, i, t, GRN)
    for r in range(P_FIRST, P_LAST + 1): common_left(ws, r, "C", "D", "E", "F")
    paste_area(ws, data.PASTE["진단"], "DIAG_ITEMS")
    widths(ws, {"A": 12, "B": 9, "C": 8, "D": 10, "E": 11, "F": 11}); ws.freeze_panes = "B3"

    # ---------------- 평가A ----------------
    ws = wb.create_sheet("평가A")
    title(ws, "A1", "평가A (1회차 공통 평가, 8문항) — 기준점수 = 평가A로 고정. 평가A가 없으면 첫 응시(평가B, 없으면 평가C)가 기준이고 그 회차는 판정 보류")
    hd = ["참가자 코드", "팀", "A 맞힌 수", "A 정답률(%)", "상태", "중복 점검", "기준점수(%)", "기준 출처",
          "기준 맞힌 수", "진단 정답률(%)", "진단−평가A(%p)", "교수 확인", "팀 편성 참고(진단+평가A %)"]
    for i, t in enumerate(hd, 1): header(ws, 2, i, t, GRN)
    for r in range(P_FIRST, P_LAST + 1):
        common_left(ws, r, "C", "D", "E", "F")
        ws[f"G{r}"] = f'=IF(A{r}="","",IF(ISNUMBER(D{r}),D{r},IF(ISNUMBER(평가B!F{r}),평가B!F{r},IF(ISNUMBER(평가C!G{r}),평가C!G{r},""))))'
        ws[f"H{r}"] = (f'=IF(A{r}="","",IF(ISNUMBER(D{r}),"평가A",IF(ISNUMBER(평가B!F{r}),"첫 응시(평가B)",'
                       f'IF(ISNUMBER(평가C!G{r}),"첫 응시(평가C)","기준 없음"))))')
        ws[f"I{r}"] = f'=IF(A{r}="","",IF(ISNUMBER(D{r}),C{r},IF(ISNUMBER(평가B!F{r}),평가B!D{r},IF(ISNUMBER(평가C!G{r}),평가C!E{r},""))))'
        ws[f"J{r}"] = f'=IF(A{r}="","",진단!D{r})'
        ws[f"K{r}"] = f'=IF(A{r}="","",IF(AND(ISNUMBER(D{r}),ISNUMBER(J{r})),ROUND(J{r}-D{r},1),""))'
        ws[f"L{r}"] = f'=IF(A{r}="","",IF(ISNUMBER(K{r}),IF(K{r}>=DIAG_GAP,"교수 확인",""),""))'
        ws[f"M{r}"] = (f'=IF(A{r}="","",IF(AND(NOT(ISNUMBER(D{r})),NOT(ISNUMBER(J{r}))),"",'
                       f'ROUND((IF(ISNUMBER(D{r}),D{r}*EVAL_ITEMS,0)+IF(ISNUMBER(J{r}),J{r}*DIAG_ITEMS,0))/'
                       f'(IF(ISNUMBER(D{r}),EVAL_ITEMS,0)+IF(ISNUMBER(J{r}),DIAG_ITEMS,0)),1)))')
    paste_area(ws, data.PASTE["평가A"], "EVAL_ITEMS")
    widths(ws, {"A": 12, "B": 9, "C": 8, "D": 10, "E": 11, "F": 11, "G": 10, "H": 15, "I": 9, "J": 10, "K": 10, "L": 10, "M": 13})
    ws.freeze_panes = "B3"

    # ---------------- 평가B / 평가C ----------------
    def eval_sheet(name, sess, first_label, hold_text, layout):
        ws = wb.create_sheet(name)
        title(ws, "A1", f"{name} ({sess}회차 공통 평가, 8문항) — 기준점수 대비 향상점수. 개인 결과는 교수님만 봄")
        for i, (t, _) in enumerate(layout, 1): header(ws, 2, i, t, GRN)
        c = {k: chr(64 + i) for i, (_, k) in enumerate(layout, 1)}
        for r in range(P_FIRST, P_LAST + 1):
            common_left(ws, r, c["cnt"], c["pct"], c["st"], c["dup"])
            ws[f"C{r}"] = f'=IF(A{r}="","",평가A!C{r})'
            if name == "평가C": ws[f"D{r}"] = f'=IF(A{r}="","",평가B!D{r})'
            P, BS, SRC, CH, PT, ST = (f"{c[k]}{r}" for k in ("pct", "base", "src", "chg", "pts", "st"))
            hold = f'{SRC}="{first_label}"'
            ws[f"{c['base']}{r}"] = f'=IF(A{r}="","",평가A!G{r})'
            ws[f"{c['src']}{r}"] = f'=IF(A{r}="","",평가A!H{r})'
            ws[f"{c['dcnt']}{r}"] = (f'=IF(A{r}="","",IF(OR(NOT(ISNUMBER({c["cnt"]}{r})),{hold},NOT(ISNUMBER(평가A!I{r}))),"",'
                                     f'ROUND({c["cnt"]}{r}-평가A!I{r},1)))')
            ws[CH] = f'=IF(A{r}="","",IF(OR({P}="",{BS}="",{hold}),"",ROUND({P}-{BS},1)))'
            ws[PT] = (f'=IF(A{r}="","",IF(OR({P}="",{BS}="",{hold}),"",IF({P}>=HIGH_CUT,PTS_HIGH,'
                      f'IF({CH}>=STEP2,PTS_HIGH,IF({CH}>=STEP1,PTS_MID,PTS_LOW)))))')
            ws[f"{c['lab']}{r}"] = (f'=IF(A{r}="","",IF({P}="","",IF({hold},"보류",IF({BS}="","",IF({P}>=HIGH_CUT,"숙달",'
                                    f'IF({CH}>=STEP2,"향상",IF({CH}>=STEP1,"유지","보완 필요")))))))')
            ws[f"{c['team']}{r}"] = f'=IF(A{r}="","",IF(AND({ST}="응시",ISNUMBER({PT})),{PT},""))'
            ws[f"{c['why']}{r}"] = (f'=IF(A{r}="","",IF({P}="","결석 - 기록 없음",IF({hold},"{hold_text}",'
                                    f'IF({BS}="","기준점수 없음",IF({P}>=HIGH_CUT,"숙달 - 이번 정답률이 숙달 기준 이상",'
                                    f'IF({CH}>=STEP2,"향상 - 기준보다 크게 오름",IF({CH}>=STEP1,"유지 - 기준과 비슷함",'
                                    f'"보완 필요 - 기준보다 크게 내려감"))))))&IF({ST}="비경쟁 응시"," / 비경쟁: 팀 점수 제외",""))')
        paste_area(ws, data.PASTE[name], "EVAL_ITEMS")
        w = {chr(64 + i): 9 for i in range(1, len(layout) + 1)}
        w.update({"A": 12, c["src"]: 15, c["lab"]: 10, c["why"]: 40, c["st"]: 11, c["dup"]: 11})
        widths(ws, w); ws.freeze_panes = "B3"
        return c

    LB = [("참가자 코드", "code"), ("팀", "team_"), ("A 맞힌 수", "a"), ("B 맞힌 수(이번)", "cnt"),
          ("변화(개, 기준 대비)", "dcnt"), ("정답률(%)", "pct"), ("상태", "st"), ("중복 점검", "dup"),
          ("기준점수(%)", "base"), ("기준 출처", "src"), ("변화(%p)", "chg"), ("향상점수", "pts"),
          ("결과 명칭", "lab"), ("팀 반영 점수", "team"), ("판정 이유", "why")]
    LC = [("참가자 코드", "code"), ("팀", "team_"), ("A 맞힌 수", "a"), ("B 맞힌 수", "b"), ("C 맞힌 수(이번)", "cnt"),
          ("변화(개, 기준 대비)", "dcnt"), ("정답률(%)", "pct"), ("상태", "st"), ("중복 점검", "dup"),
          ("기준점수(%)", "base"), ("기준 출처", "src"), ("변화(%p)", "chg"), ("향상점수", "pts"),
          ("결과 명칭", "lab"), ("팀 반영 점수", "team"), ("판정 이유", "why")]
    cB = eval_sheet("평가B", 2, "첫 응시(평가B)", "판정 보류 - 평가A가 없어 이번 결과가 기준점수(평가C부터 판정)", LB)
    cC = eval_sheet("평가C", 3, "첫 응시(평가C)", "판정 보류 - 평가A·B가 없어 이번 결과가 기준점수", LC)

    # ---------------- 팀 점수 ----------------
    ws = wb.create_sheet("팀 점수")
    title(ws, "A1", "공개용 (활동이 끝난 뒤에만) — 팀 점수 = 판정된 팀원의 향상점수 평균. 결석·비경쟁·판정 보류 제외. 15점 이상 '학습목표 달성'")
    ws["L1"] = "교수용 — 화면에 띄우지 마세요"; ws["L1"].font = Font(bold=True)
    for i, t in enumerate(["팀", "2회차(평가B) 팀 점수", "2회차 판정", "2회차 순위", "3회차(평가C) 팀 점수", "3회차 판정",
                           "3회차 순위", "종합 평균", "종합 순위(상위 3)", "챌린지 완주"], 1): header(ws, 2, i, t, BLU)
    for i, t in enumerate(["팀원 수", "평가B 반영 인원", "평가C 반영 인원"], 12): header(ws, 2, i, t, PNK)
    for r in range(3, 13):
        ws[f"A{r}"].fill = fill(YEL)
        if r - 3 < len(data.TEAMS): ws[f"A{r}"] = data.TEAMS[r - 3]
        for sc, cc, rc, sh, cm, cnt in (("B", "C", "D", "평가B", cB, "M"), ("E", "F", "G", "평가C", cC, "N")):
            ws[f"{sc}{r}"] = f'=IF(A{r}="","",IF({cnt}{r}=0,"",ROUND(SUMIFS({sh}!${cm["team"]}$3:${cm["team"]}$52,{sh}!$B$3:$B$52,A{r})/{cnt}{r},1)))'
            ws[f"{cc}{r}"] = f'=IF(A{r}="","",IF({sc}{r}="","기록 없음",IF({sc}{r}>=GOAL,"학습목표 달성","도전 중")))'
            ws[f"{rc}{r}"] = f'=IF(A{r}="","",IF({sc}{r}="","",IF({sc}{r}<GOAL,"",RANK({sc}{r},${sc}$3:${sc}$12))))'
            ws[f"{cnt}{r}"] = f'=IF(A{r}="","",COUNTIFS({sh}!$B$3:$B$52,A{r},{sh}!${cm["team"]}$3:${cm["team"]}$52,">=0"))'
        ws[f"H{r}"] = f'=IF(A{r}="","",IF(COUNT(B{r},E{r})=0,"",ROUND(AVERAGE(B{r},E{r}),1)))'
        ws[f"I{r}"] = f'=IF(OR(A{r}="",H{r}=""),"",IF(RANK(H{r},$H$3:$H$12)<=3,RANK(H{r},$H$3:$H$12),""))'
        ws[f"J{r}"] = (f'=IF(A{r}="","",IF(COUNT(B{r},E{r})=0,"",IF(AND(OR(NOT(ISNUMBER(B{r})),B{r}>=GOAL),'
                       f'OR(NOT(ISNUMBER(E{r})),E{r}>=GOAL)),"챌린지 완주","")))')
        ws[f"L{r}"] = f'=IF(A{r}="","",COUNTIF(평가A!$B$3:$B$52,A{r}))'
    title(ws, "A15", "반 전체 요약 (익명 공개용 + 합격 기준 점검)")
    for i, t in enumerate(["시트", "명단 인원", "출석 인원(직접 입력)", "응시 인원(비경쟁 포함)", "완료율(%)", "평균 정답률(%)",
                           "중복 코드", "명단에 없는 코드", "출석했는데 기록 없음", "비경쟁 응시"], 1): header(ws, 16, i, t, BLU)
    cols = {"진단": ("D", "E", "F"), "평가A": ("D", "E", "F"), "평가B": (cB["pct"], cB["st"], cB["dup"]), "평가C": (cC["pct"], cC["st"], cC["dup"])}
    for r, (sh, (pc, st, dp)) in zip(range(17, 21), cols.items()):
        ws[f"A{r}"] = sh; ws[f"C{r}"].fill = fill(YEL)
        ws[f"B{r}"] = '=COUNTIF(참가자!$A$3:$A$52,"?*")'
        ws[f"D{r}"] = f'=COUNTIF({sh}!${st}$3:${st}$52,"*응시")'
        ws[f"E{r}"] = f'=IF(IF(C{r}="",B{r},C{r})=0,"",ROUND(D{r}/IF(C{r}="",B{r},C{r})*100,1))'
        ws[f"F{r}"] = f'=IF(D{r}=0,"",ROUND(AVERAGE({sh}!${pc}$3:${pc}$52),1))'
        ws[f"G{r}"] = f'=COUNTIF({sh}!${dp}$3:${dp}$52,"중복*")'
        ws[f"H{r}"] = f'=COUNTIF({sh}!$Z$3:$Z$102,"명단*")'
        ws[f"I{r}"] = f'=IF(C{r}="","(출석 인원 입력 시)",MAX(0,C{r}-D{r}))'
        ws[f"J{r}"] = f'=COUNTIF({sh}!${st}$3:${st}$52,"비경쟁*")'
    title(ws, "A22", "기준점수 출처 인원 (숫자만 — 누구인지는 적지 않음)")
    for r, (lab, f) in zip(range(23, 28), [
            ("기준 = 평가A", '=COUNTIF(평가A!$H$3:$H$52,"평가A")'),
            ("기준 = 첫 응시(평가B) → 평가B 판정 보류", '=COUNTIF(평가A!$H$3:$H$52,"첫 응시(평가B)")'),
            ("기준 = 첫 응시(평가C) → 평가C 판정 보류", '=COUNTIF(평가A!$H$3:$H$52,"첫 응시(평가C)")'),
            ("기준 없음(평가A·B·C 모두 없음)", '=COUNTIF(평가A!$H$3:$H$52,"기준 없음")'),
            ("'교수 확인' 표시(진단이 평가A보다 기준값 이상 높음)", '=COUNTIF(평가A!$L$3:$L$52,"교수 확인")')]):
        ws[f"A{r}"] = lab; ws[f"D{r}"] = f
    title(ws, "A29", "개인 결과 명칭 인원 (반 전체 숫자만, 비경쟁 응시 포함)")
    for i, t in enumerate(["회차", "숙달", "향상", "유지", "보완 필요", "판정 보류", "결석"], 1): header(ws, 30, i, t, BLU)
    for r, (sh, cm) in zip((31, 32), (("평가B", cB), ("평가C", cC))):
        ws[f"A{r}"] = sh
        for col, lab in zip("BCDEF", ["숙달", "향상", "유지", "보완 필요", "보류"]):
            ws[f"{col}{r}"] = f'=COUNTIF({sh}!${cm["lab"]}$3:${cm["lab"]}$52,"{lab}")'
        ws[f"G{r}"] = f'=COUNTIF({sh}!${cm["st"]}$3:${cm["st"]}$52,"결석")'
    ws["A34"] = ("완료율 = 응시 인원 / 출석 인원(비어 있으면 명단 인원). 합격 기준: 완료율 90% 이상, 중복·명단 외 코드·출석했는데 기록 없음 모두 0건. "
                 "진단은 1회차 같은 날이라 완료율은 평가A 기준으로 봅니다. 이 숫자들은 탐색적 파일럿의 참여 촉진 지표이며, 학습 성장을 엄밀하게 잰 값이 아닙니다.")
    widths(ws, {"A": 14, "B": 12, "C": 12, "D": 11, "E": 12, "F": 12, "G": 10, "H": 12, "I": 13, "J": 12, "K": 3, "L": 9, "M": 12, "N": 12})

    # ---------------- 사용법 ----------------
    ws = wb.create_sheet("사용법")
    title(ws, "A1", "[2026-11-02] 결재2 ChatGPT 교차검토 반영: 결과 명칭(숙달·향상/유지/보완 필요), 팀 15점 '학습목표 달성', 기준점수 = 평가A 고정(진단 보정·반 평균 대입 삭제)")
    ws["A2"] = "값이 다시 바뀌면 아래 노란 칸만 고치면 모든 시트가 자동으로 바뀝니다. 시범 기간 중에는 바꾸지 않습니다. 이 결과는 탐색적 파일럿의 '참여 촉진 지표'입니다(엄밀한 성장 측정 아님)."
    ws["A3"] = VERIFY_NOTE
    title(ws, "A4", "기준값 (노란 칸만 수정)")
    for i, t in enumerate(["이름표", "설명", "값"], 1): header(ws, 5, i, t, GRY)
    params = [
        ("HIGH_CUT", "숙달 기준(%) — 이번 정답률이 이 값 이상이면 기준과 상관없이 최고점, 결과 명칭 '숙달' (8문항 중 7개 = 87.5%)"),
        ("STEP1", "0점 경계(%p) — 기준 대비 변화가 이 값 이상이면 '유지', 미만이면 '보완 필요' (−20%p: 2문제 이상 덜 맞히면 0점, 정확히 −20%p는 유지)"),
        ("STEP2", "20점 경계(%p) — 기준 대비 변화가 이 값 이상이면 최고점, 결과 명칭 '향상' (+20%p: 2문제 이상 더 맞히면)"),
        ("PTS_LOW", "'보완 필요'일 때 향상점수"),
        ("PTS_MID", "'유지'일 때 향상점수"),
        ("PTS_HIGH", "'숙달' 또는 '향상'일 때 향상점수"),
        ("GOAL", "팀 점수가 이 값 이상이면 '학습목표 달성', 미만이면 '도전 중' (회차 순위는 달성 팀만)"),
        ("DIAG_GAP", "교수 확인 기준(%p, 임시값) — 진단 정답률 − 평가A 정답률이 이 값 이상이면(진단이 더 높을 때만) 평가A 시트에 '교수 확인' 표시만 함(점수는 그대로)"),
        ("DIAG_ITEMS", "진단 문항 수 — '팀 편성 참고' 합산 정답률 계산용"),
        ("EVAL_ITEMS", "평가 A·B·C 문항 수 — '팀 편성 참고' 계산, 정답률만 붙여넣었을 때 맞힌 수 환산용"),
    ]
    for r, (nm, desc) in enumerate(params, 6):
        ws[f"A{r}"] = nm; ws[f"B{r}"] = desc; ws[f"C{r}"] = vals[nm]; ws[f"C{r}"].fill = fill(YEL)
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"'사용법'!$C${r}")
    rules = [
        ("A18", "한눈에 보는 규칙 (현재 노란 칸 값 기준)", True),
        ("A19", "1. 1회차 = 진단(12문항, 팀 편성·수준 파악) + 평가A(8문항, 기준점수). 2회차 = 평가B, 3회차 = 평가C (각 8문항, 1문항 = 12.5%p). 팀 챌린지는 2·3회차.", False),
        ("A20", "2. 기준점수 = 평가A 정답률로 고정해 평가B·C에 똑같이 씁니다. 진단은 기준점수에 쓰지 않습니다(진단 평균 보정 없음).", False),
        ("A21", "   평가A가 없으면(1회차 결석·일찍 나감·새 참가자) 첫 응시(평가B, 그것도 없으면 평가C)가 기준점수가 되고, 그 회차는 '판정 보류'(점수 비움, 팀 평균에서 제외). 다음 회차부터 판정합니다.", False),
        ("A22", "   진단이 평가A보다 교수 확인 기준(기본 30%p, 임시값) 이상 높으면 '교수 확인' 표시만 합니다(일부러 평가A를 낮게 봤을 가능성 참고용). 점수는 바꾸지 않습니다.", False),
        ("A23", "3. 개인 결과 (위에서부터 먼저 맞는 것 하나): 이번 정답률 80% 이상(8개 중 7개) → 20점 '숙달' / 기준 +20%p 이상 → 20점 '향상' / 기준 −20%p 이상 → 10점 '유지' / 그보다 낮음 → 0점 '보완 필요'", False),
        ("A24", "   한 문제 차이(±12.5%p)는 위아래 모두 '유지'. 경계값은 포함: 정확히 80% → 숙달, 정확히 +20%p → 향상, 정확히 −20%p → 유지(10점).", False),
        ("A25", "4. 팀 점수 = 그 회차 '응시'하고 판정된 팀원의 향상점수 평균(소수 첫째 자리 반올림). 결석·비경쟁·명단에 없는 코드·판정 보류는 빠집니다. 0명이면 '기록 없음'.", False),
        ("A26", "5. 팀 점수 15점 이상 '학습목표 달성', 미만 '도전 중'. 회차 순위는 달성 팀만, 같은 점수는 같은 순위. 종합 순위는 기록 있는 회차 평균으로 상위 3위까지만.", False),
        ("A27", "6. 챌린지 완주 = 2·3회차 모두 학습목표 달성. 한 회차가 '기록 없음'이면 남은 1회만으로 판정.", False),
        ("A28", "7. 정답률은 소수 첫째 자리에서 반올림해 저장하고, 이후 계산은 그 값으로 합니다(예: 7/12 = 58.3%). 원점수(맞힌 수) A·B·C와 변화(개)를 함께 보여 줍니다.", False),
        ("A29", "8. 팀 편성 참고: 평가A 시트 M열 = 진단 + 평가A 20문항 합산 정답률. 높은 순으로 줄 세워 지그재그로 팀을 짭니다.", False),
        ("A30", "9. 해석 주의: 4지선다라 찍기 운이 큽니다(모두 찍어도 평가B에서 20점이 나올 확률 약 19%, 가정에 따른 계산). 한 번의 20점을 '실력이 늘었다'로 말하지 않습니다.", False),
        ("A32", "사용 순서", True),
        ("A33", "1단계 — [참가자] 시트: 노란 칸에 6자리 참가자 코드와 팀을 적습니다. 실명·학번·연락처는 절대 적지 마세요.", False),
        ("A34", "2단계 — 각 평가 시트의 오른쪽 노란 '붙여넣기 영역'(R~U열, 3행부터)에 결과를 붙여넣습니다. 왼쪽 초록 칸과 회색 칸은 자동 계산이니 손대지 마세요.", False),
        ("A35", "   1회차: 진단 12문항 → [진단], 평가A 8문항 → [평가A]. 2회차 → [평가B], 3회차 → [평가C]. (연습 문항 결과는 넣지 않습니다.)", False),
        ("A36", "3단계 — [팀 점수] 시트: 팀 이름(노란 칸)을 확인하고, 반 전체 요약에 그날 출석 인원을 적습니다. 화면에는 왼쪽 '공개용' 표와 반 전체 요약만 띄웁니다. 오른쪽 분홍 '교수용' 칸은 띄우지 마세요.", False),
        ("A38", "붙여넣기 방법 (Socrative 또는 우리 웹앱 CSV)", True),
        ("A39", "R열 = 참가자 코드 / S열 = 맞힌 문항 수 / T열 = 전체 문항 수 / U열 = 정답률(%) — S·T가 있으면 U는 비워 두고, 정답률만 있으면 S·T를 비우고 U에만 붙여넣습니다(U에 값이 있으면 U를 우선 사용, 맞힌 수는 정답률 × 문항 수로 환산해 소수가 나올 수 있음).", False),
        ("A40", "Socrative: 학생 '이름' 칸에 참가자 코드를 입력하게 안내 → 끝난 뒤 보고서의 이름 열과 점수 열만 복사해 붙여넣기. (무료 요금제 내보내기 형식은 미확인)", False),
        ("A41", "우리 웹앱(트랙 B): 내보내기 CSV 열 순서를 '참가자코드, 맞힌 수, 문항 수'로 맞추면 R열에 그대로 붙여넣을 수 있습니다. 코드 앞뒤 빈칸·대소문자는 자동으로 정리됩니다.", False),
        ("A42", "붙여넣은 표에 실명이나 학번 열이 섞여 있으면 붙여넣기 전에 그 열을 지우세요.", False),
        ("A44", "비경쟁·예외 처리", True),
        ("A45", "코드 뒤 '-N'을 붙인 응시(예: K7M2QX-N)는 '비경쟁 응시'. 점수·명칭은 계산하지만 팀 평균·순위에는 넣지 않습니다. 평가 시작 3분 뒤 도착자는 R열 코드 뒤에 '-N'을 직접 붙여 주세요.", False),
        ("A46", "명단에 없는 코드(오타)는 계산에 들어가지 않습니다('명단에 없는 코드' 표시, 기록 점검에서 1건). '중복! 확인' = 같은 코드가 두 번 → 위에 있는 줄이 쓰입니다(먼저 끝까지 제출한 것이 위에 오게).", False),
        ("A47", "팀 출석이 1명뿐인 회차(다른 팀 임시 합류)는 자동 처리가 없습니다. 교수님 판단으로 메모해 두세요.", False),
        ("A49", "개인정보·공개 원칙", True),
        ("A50", "개인별 정답률·향상점수·결과 명칭(진단~평가C 시트)은 교수님만 봅니다. 팀 순위는 활동이 끝난 뒤에만 공개하고, 팀 평균을 몇 명으로 나눴는지는 보여 주지 않습니다.", False),
        ("A51", "이 파일의 12명 데이터는 모두 가짜 예시입니다. 실제 사용 전 참가자·붙여넣기 영역의 예시 값을 지우세요. 학기 종료 후 파일을 삭제합니다.", False),
    ]
    for cell, text, bold in rules:
        ws[cell] = text
        if bold: ws[cell].font = Font(bold=True)
    for r in range(6, 16): ws[f"B{r}"].alignment = WRAP
    widths(ws, {"A": 14, "B": 90, "C": 10})


    wb.calculation.fullCalcOnLoad = True  # 엑셀·구글 시트에서 열 때 다시 계산
    wb.save(path)

VERIFY_NOTE = ("[검증 기록 2026-11-02] 같은 규칙을 파이썬으로 따로 계산해 LibreOffice 재계산 값과 대조: 확정값 + 노란 칸 변경 2가지 = "
               "3가지 경우 × 601개 항목 모두 일치(총 1,803개), 오류 셀 0개. 자세한 내용: 05_성장팀/2026-11-01_팀챌린지_설문과합격기준.md 5절")

if __name__ == "__main__":
    out = sys.argv[1]
    sc = sys.argv[2] if len(sys.argv) > 2 else "1_확정값"
    build(out, data.SCENARIOS[sc])
    print("saved", out, sc)
